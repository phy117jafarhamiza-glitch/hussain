# -*- coding: utf-8 -*-
"""
توليد تقرير Word (.docx) لكل لاعب يلخّص تطوره عبر برنامج التأهيل (8 أسابيع)،
متضمناً: بيانات اللاعب، الاختبارات القبلية/البعدية، نتائج تحليل EMG
بالشبكة العصبية، وتوصيات وحدة التحكم الذكية (المحاكاة).
"""
import os
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from utils.constants import TEST_NAMES

REPORTS_DIR = "data/reports"
os.makedirs(REPORTS_DIR, exist_ok=True)


def _set_rtl(paragraph):
    """يضبط اتجاه الفقرة من اليمين لليسار (للنص العربي)."""
    pPr = paragraph._p.get_or_add_pPr()
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    bidi = OxmlElement("w:bidi")
    pPr.append(bidi)


def generate_subject_report(subject: dict, measurements: list, emg_analyses: list,
                             controller_logs: list, out_path: str = None) -> str:
    doc = Document()

    title = doc.add_heading(f"تقرير تطور اللاعب: {subject['name']}", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_rtl(title)

    meta = doc.add_paragraph(
        f"المجموعة: {subject['group_name']}   |   العمر: {subject.get('age') or '-'}   |   "
        f"الرياضة: {subject.get('sport') or '-'}\n"
        f"تاريخ إصدار التقرير: {datetime.now().strftime('%Y-%m-%d')}"
    )
    _set_rtl(meta)

    doc.add_heading("١. نتائج الاختبارات عبر مراحل البرنامج", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "المرحلة"
    hdr[1].text = "الاختبار"
    hdr[2].text = "القيمة"
    for m in measurements:
        row = table.add_row().cells
        row[0].text = str(m.get("stage", ""))
        row[1].text = TEST_NAMES.get(m.get("test_key"), m.get("test_key", ""))
        row[2].text = "" if m.get("value") is None else str(m.get("value"))

    doc.add_heading("٢. نتائج تحليل إشارة EMG بالشبكة العصبية", level=2)
    if emg_analyses:
        t2 = doc.add_table(rows=1, cols=4)
        t2.style = "Light Grid Accent 1"
        h = t2.rows[0].cells
        h[0].text = "المرحلة"
        h[1].text = "العضلة"
        h[2].text = "التصنيف"
        h[3].text = "الثقة"
        for e in emg_analyses:
            row = t2.add_row().cells
            row[0].text = str(e.get("stage", ""))
            row[1].text = str(e.get("muscle", ""))
            row[2].text = str(e.get("predicted_label", ""))
            row[3].text = f"{e.get('confidence', 0):.2f}"
    else:
        doc.add_paragraph("لا توجد بيانات تحليل EMG مسجلة لهذا اللاعب بعد.")

    doc.add_heading("٣. سجل توصيات الجهاز الذكي (محاكاة وحدة التحكم بالذكاء الاصطناعي)", level=2)
    if controller_logs:
        t3 = doc.add_table(rows=1, cols=4)
        t3.style = "Light Grid Accent 1"
        h = t3.rows[0].cells
        h[0].text = "الاتجاه"
        h[1].text = "شدة EMG (%)"
        h[2].text = "زاوية الرأس"
        h[3].text = "المقاومة الموصى بها"
        for c in controller_logs:
            row = t3.add_row().cells
            row[0].text = str(c.get("direction", ""))
            row[1].text = str(c.get("emg_activation", ""))
            row[2].text = str(c.get("head_angle", ""))
            row[3].text = f"{c.get('recommended_resistance','')} ({c.get('recommended_resistance_value','')}%)"
    else:
        doc.add_paragraph("لا توجد سجلات توصيات مسجلة لهذا اللاعب بعد.")

    note = doc.add_paragraph(
        "ملاحظة: هذا التقرير أداة متابعة برمجية مساندة للأطروحة، ولا يغني عن "
        "التقييم السريري المباشر من قبل المشرف والباحث."
    )
    note.runs[0].font.size = Pt(9)
    note.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    _set_rtl(note)

    if out_path is None:
        safe_name = "".join(c for c in subject["name"] if c.isalnum() or c in " _-")
        out_path = os.path.join(REPORTS_DIR, f"تقرير_{safe_name}_{subject['id']}.docx")
    doc.save(out_path)
    return out_path
